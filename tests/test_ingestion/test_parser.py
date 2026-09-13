"""文档解析器单元测试。"""

import base64
import io
import zlib
from pathlib import Path

from app.ingestion.parser import DocumentParser


def _build_pdf(pages: list[str], *, image_page: int | None = None) -> bytes:
    """手工构造最小 PDF，免引入 PDF 生成依赖。

    pages: 每页的内容流（含 BT/ET 文本指令或 re 矩形指令）。
    image_page: 在该页（0-based）挂一个 1x1 红色图片并绘制，模拟扫描件。
    """
    n = len(pages)
    total = 4 + 2 * n + (1 if image_page is not None else 0)
    bodies: list[bytes] = [b""] * total

    def obj(num: int, body: bytes) -> None:
        bodies[num - 1] = body

    kids = " ".join(f"{4 + i * 2} 0 R" for i in range(n))
    obj(1, b"<< /Type /Catalog /Pages 2 0 R >>")
    obj(2, f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    obj(3, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    xobj_ref = ""
    if image_page is not None:
        img_num = total
        raw = zlib.compress(b"\xff\x00\x00")
        obj(img_num, (
            f"<< /Type /XObject /Subtype /Image /Width 1 /Height 1 "
            f"/ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /FlateDecode "
            f"/Length {len(raw)} >>\nstream\n".encode()
            + raw + b"\nendstream"
        ))
        xobj_ref = f" /XObject << /Im1 {img_num} 0 R >>"

    for i, content in enumerate(pages):
        page_num, content_num = 4 + i * 2, 5 + i * 2
        extra = xobj_ref if i == image_page else ""
        obj(page_num, (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >>{extra} >> "
            f"/Contents {content_num} 0 R >>"
        ).encode())
        data = content.encode("latin-1")
        obj(content_num, (
            f"<< /Length {len(data)} >>\nstream\n".encode() + data + b"\nendstream"
        ))

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for num, body in enumerate(bodies, 1):
        offsets.append(len(out))
        out += f"{num} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {total + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    out += "".join(f"{off:010d} 00000 n \n" for off in offsets).encode()
    out += (
        f"trailer\n<< /Size {total + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n"
    ).encode()
    return bytes(out)


def test_parse_markdown(tmp_path: Path):
    md = tmp_path / "test.md"
    md.write_text("# 标题\n\n这是内容。", encoding="utf-8")

    result = DocumentParser().parse(md)

    assert result.filename == "test.md"
    assert result.file_type == ".md"
    assert len(result.pages) == 1
    assert "标题" in result.pages[0]
    assert "内容" in result.pages[0]


def test_parse_python_code(tmp_path: Path):
    py = tmp_path / "auth.py"
    py.write_text("def login():\n    pass\n", encoding="utf-8")

    result = DocumentParser().parse(py)

    assert result.file_type == ".py"
    assert "def login" in result.pages[0]


def test_parse_txt(tmp_path: Path):
    txt = tmp_path / "notes.txt"
    txt.write_text("第一条笔记\n第二条笔记\n", encoding="utf-8")

    result = DocumentParser().parse(txt)

    assert result.file_type == ".txt"
    assert "第一条笔记" in result.pages[0]


def test_unsupported_format_raises(tmp_path: Path):
    fake = tmp_path / "data.csv"
    fake.write_text("a,b,c\n1,2,3", encoding="utf-8")

    import pytest
    with pytest.raises(ValueError, match="不支持"):
        DocumentParser().parse(fake)


def test_missing_file_raises():
    import pytest
    with pytest.raises(FileNotFoundError):
        DocumentParser().parse("nonexistent_file.xyz")


def test_parse_docx(tmp_path: Path):
    """测试 Word 文档解析（用 python-docx 生成测试文件）。"""
    from docx import Document as DocxDocument

    docx_path = tmp_path / "test.docx"
    doc = DocxDocument()
    doc.add_paragraph("第一段内容")
    doc.add_paragraph("第二段内容")
    doc.save(str(docx_path))

    result = DocumentParser().parse(docx_path)

    assert result.file_type == ".docx"
    assert "第一段内容" in result.pages[0]
    assert "第二段内容" in result.pages[0]


def test_parse_docx_table_order_nested_and_merge(tmp_path: Path):
    """表格与段落按文档顺序交错；嵌套表提取；合并单元格不重复。"""
    from docx import Document as DocxDocument

    docx_path = tmp_path / "tables.docx"
    doc = DocxDocument()
    doc.add_paragraph("表格前说明")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "项目"
    t.cell(0, 1).text = "进度"
    t.cell(1, 0).text = "摄入管线"
    t.cell(1, 1).text = "已完成"
    doc.add_paragraph("表格后结论")

    outer = doc.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "外层单元格"
    inner = outer.cell(0, 0).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "内A"
    inner.cell(0, 1).text = "内B"

    merged = doc.add_table(rows=1, cols=3)
    merged.cell(0, 0).merge(merged.cell(0, 1))
    merged.cell(0, 0).text = "合并表头"
    merged.cell(0, 2).text = "备注"
    doc.save(str(docx_path))

    result = DocumentParser().parse(docx_path)
    text = result.pages[0]

    # 阅读顺序：段落与表格按文档实际顺序交错
    assert text.index("表格前说明") < text.index("项目 | 进度") < text.index("表格后结论")
    # 嵌套表不再静默丢失
    assert "外层单元格" in text
    assert "内A | 内B" in text
    # 合并单元格按底层元素去重
    assert "合并表头 | 备注" in text
    assert "合并表头 | 合并表头" not in text


def test_parse_xlsx(tmp_path: Path):
    """测试 Excel 解析。"""
    from openpyxl import Workbook

    xlsx_path = tmp_path / "test.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["姓名", "年龄"])
    ws.append(["张三", 30])
    wb.save(str(xlsx_path))

    result = DocumentParser().parse(xlsx_path)

    assert result.file_type == ".xlsx"
    assert "张三" in result.pages[0]
    assert "姓名" in result.pages[0]


def test_parse_pptx_includes_table_and_text(tmp_path: Path):
    """PPT 表格此前因 GraphicFrame 无 .text 属性被静默丢弃，回归验证。"""
    from pptx import Presentation
    from pptx.util import Inches

    pptx_path = tmp_path / "report.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式
    box = slide.shapes.add_textbox(Inches(1), Inches(0.2), Inches(4), Inches(0.5))
    box.text_frame.text = "季度汇报"
    frame = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
    table = frame.table
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "进度"
    table.cell(1, 0).text = "摄入管线"
    table.cell(1, 1).text = "已完成"
    prs.save(str(pptx_path))

    result = DocumentParser().parse(pptx_path)

    text = result.pages[0]
    assert "季度汇报" in text
    assert "项目 | 进度" in text
    assert "摄入管线 | 已完成" in text


def test_parse_xlsx_sheet_prefix_and_header_alignment(tmp_path: Path):
    """多 sheet 时 sheet 名进每行前缀；数据行对齐成"列名: 值"。"""
    from openpyxl import Workbook

    xlsx_path = tmp_path / "sales.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "员工"
    ws.append(["姓名", "年龄"])
    ws.append(["张三", 30])
    ws2 = wb.create_sheet("销售")
    ws2.append(["产品", "数量"])
    ws2.append(["键盘", 12])
    wb.save(str(xlsx_path))

    result = DocumentParser().parse(xlsx_path)
    text = result.pages[0]

    assert "[工作表 员工] 表头: 姓名 | 年龄" in text
    assert "[工作表 员工] 姓名: 张三 | 年龄: 30" in text
    assert "[工作表 销售] 表头: 产品 | 数量" in text
    assert "[工作表 销售] 产品: 键盘 | 数量: 12" in text


def test_parse_xlsx_sparse_columns_keep_position(tmp_path: Path):
    """行中间有空单元格时按列位置对齐表头，不错位。"""
    from openpyxl import Workbook

    xlsx_path = tmp_path / "sparse.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["名称", "备注", "数量"])
    ws.append(["螺丝", None, 100])  # 中间列为空
    wb.save(str(xlsx_path))

    result = DocumentParser().parse(xlsx_path)
    text = result.pages[0]

    assert "[工作表 数据] 名称: 螺丝 | 数量: 100" in text


_PDF_TEXT_PAGE = "BT /F1 12 Tf 72 720 Td (ZhiYuan knowledge base) Tj ET"
# 2x2 表格：4 个描边矩形 + 每格一行文本，pdfplumber 按线框识别
_PDF_TABLE_PAGE = (
    "q "
    "30 700 150 24 re S 180 700 150 24 re S "
    "30 676 150 24 re S 180 676 150 24 re S "
    "BT /F1 12 Tf 40 708 Td (Item) Tj ET "
    "BT /F1 12 Tf 190 708 Td (Status) Tj ET "
    "BT /F1 12 Tf 40 684 Td (Ingest) Tj ET "
    "BT /F1 12 Tf 190 684 Td (Done) Tj ET "
    "Q"
)


def test_parse_pdf_extracts_text_and_table(tmp_path: Path):
    """PDF 文本正常提取；表格按线框识别为"单元格 | 单元格"结构化行。"""
    pdf_path = tmp_path / "doc.pdf"
    pdf_path.write_bytes(_build_pdf([_PDF_TEXT_PAGE, _PDF_TABLE_PAGE]))

    result = DocumentParser().parse(pdf_path)

    assert result.file_type == ".pdf"
    assert len(result.pages) == 2
    assert "ZhiYuan knowledge base" in result.pages[0]
    assert "Item | Status" in result.pages[1]
    assert "Ingest | Done" in result.pages[1]
    assert result.image_only_pages == []


def test_parse_pdf_flags_image_only_pages(tmp_path: Path):
    """有图无文的页（扫描件）进 image_only_pages，页面占位空串保持索引对齐。"""
    pdf_path = tmp_path / "scan.pdf"
    image_draw = "q 200 0 0 200 100 400 cm /Im1 Do Q"
    pdf_path.write_bytes(_build_pdf([_PDF_TEXT_PAGE, image_draw], image_page=1))

    result = DocumentParser().parse(pdf_path)

    assert len(result.pages) == 2
    assert "ZhiYuan knowledge base" in result.pages[0]
    assert result.pages[1] == ""
    assert result.image_only_pages == [2]


def test_parse_pptx_flags_picture_only_slide(tmp_path: Path):
    """只有图片的 PPT 页进 image_only_pages，占位空串保持索引对齐。"""
    from pptx import Presentation
    from pptx.util import Inches

    png_1px = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
        "AAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
    )
    pptx_path = tmp_path / "imgs.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式
    slide.shapes.add_picture(
        io.BytesIO(png_1px), Inches(1), Inches(1), Inches(2), Inches(2)
    )
    prs.save(str(pptx_path))

    result = DocumentParser().parse(pptx_path)

    assert result.pages == [""]
    assert result.image_only_pages == [1]
