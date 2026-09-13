"""把文档中的图片页取出来：PDF 整页栅格化，PPT 抽取内嵌位图。

只负责取图，不负责识别（识别在 app.llm.vision）。总量受 max_requests
约束——每个请求都是一次 VLM API 调用，超限的页保留"未提取"标注。
"""

from __future__ import annotations

import io
from pathlib import Path

# 144 DPI：扫描件文字 OCR 足够清晰，JPEG 后体积可控
_RENDER_SCALE = 2.0
RASTER_MIMES = {"image/png", "image/jpeg", "image/webp", "image/gif"}


def collect_page_images(
    path: str | Path,
    page_numbers: list[int],
    max_requests: int,
) -> dict[int, list[tuple[bytes, str]]]:
    """取每页的图片负载，返回 {页码: [(图片字节, mime), ...]}。

    - PDF：有图无文的整页栅格化为 JPEG（扫描件）
    - PPTX：抽取该页内嵌位图（png/jpeg/webp/gif，矢量格式跳过）
    """
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return _pdf_page_images(path, page_numbers, max_requests)
    if ext == ".pptx":
        return _pptx_page_images(path, page_numbers, max_requests)
    return {}


def _pdf_page_images(
    path: str | Path,
    page_numbers: list[int],
    max_requests: int,
) -> dict[int, list[tuple[bytes, str]]]:
    import pypdfium2 as pdfium

    out: dict[int, list[tuple[bytes, str]]] = {}
    budget = max_requests
    with pdfium.PdfDocument(str(path)) as doc:
        for page_no in page_numbers:
            if budget <= 0:
                break
            if page_no < 1 or page_no > len(doc):
                continue
            bitmap = doc[page_no - 1].render(scale=_RENDER_SCALE)
            buf = io.BytesIO()
            bitmap.to_pil().convert("RGB").save(buf, "JPEG", quality=85)
            out[page_no] = [(buf.getvalue(), "image/jpeg")]
            budget -= 1
    return out


def _pptx_page_images(
    path: str | Path,
    page_numbers: list[int],
    max_requests: int,
) -> dict[int, list[tuple[bytes, str]]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    slides = list(prs.slides)
    out: dict[int, list[tuple[bytes, str]]] = {}
    budget = max_requests

    for page_no in page_numbers:
        if budget <= 0:
            break
        if page_no < 1 or page_no > len(slides):
            continue
        pictures: list[tuple[bytes, str]] = []
        for shape in slides[page_no - 1].shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            if budget <= 0:
                break
            image = shape.image
            if image.content_type not in RASTER_MIMES:
                continue
            pictures.append((image.blob, image.content_type))
            budget -= 1
        if pictures:
            out[page_no] = pictures
    return out
